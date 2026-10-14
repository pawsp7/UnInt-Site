#!/usr/bin/env python3
"""Editable Word copy of the branded event proposal template."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

EVENTS = Path(__file__).resolve().parent
REPO = EVENTS.parent
CHARTER = REPO / "charter"
sys.path.insert(0, str(CHARTER))

import revamp_copy as copy
from revamp_kit import ASSETS, export as export_kit

OUTPUT = EVENTS / "event-proposal-TEMPLATE.docx"
FONT = "Inter"

INK = RGBColor(0x03, 0x02, 0x05)
WINE = RGBColor(0x59, 0x13, 0x37)
SLATE = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIME_HEX = "E8F6C4"
PINK_HEX = "FAD6E8"
INK_HEX = "030205"
PAPER_HEX = "F9F9FA"

POSTER_QTY = 25
POSTER_COST = 12.00
NAMETAG_QTY = 30
NAMETAG_COST = 3.00
PRINTING_COST = POSTER_COST + NAMETAG_COST
MISC_COST = 5.00
BUILT_IN_TOTAL = PRINTING_COST + MISC_COST
PRINTING_DETAILS = (
    f"{POSTER_QTY}× tabloid 11×17 posters (${POSTER_COST:.0f}) and "
    f"{NAMETAG_QTY} name tags (${NAMETAG_COST:.0f}). Unaltered Saitsa Clubs "
    "logo, saitsa.com/Poster-Submission"
)


def money(value: float) -> str:
    return f"${value:.2f}"


def set_cell_margins(cell, *, twips=120) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_width(cell, twips: int) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(twips))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, *, size=11, bold=False, italic=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def style_paragraph(p, *, space_before=0, space_after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    return run


def heading(doc, text, level=1, center=False):
    p = doc.add_heading(text, level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(
            run,
            size=22 if level == 1 else 14 if level == 2 else 12,
            bold=True,
            color=INK,
        )
    style_paragraph(p, space_before=12 if level > 1 else 6, space_after=8)
    return p


def body(doc, text, *, italic=False, bold=False, color=INK, after=8, size=11, center=False):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=after)
    add_text(p, text, size=size, italic=italic, bold=bold, color=color)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Inches(0.25)
    style_paragraph(p, space_before=1, space_after=3, line=1.12)
    add_text(p, text, size=11)
    return p


def picture(doc, name, width, center=True):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=4, space_after=6)
    p.add_run().add_picture(str(ASSETS / name), width=Inches(width))
    return p


def signature(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=18, space_after=2)
    add_text(p, "Primary Event Organizer — [NAME AND EXECUTIVE ROLE]", size=10, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_after=8)
    add_text(
        p,
        "Signature: ____________________     Date: ______________",
        size=10,
        color=SLATE,
    )


def set_table_borders(table, color="D7D7DA"):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)
        borders.append(node)
    tblPr.append(borders)


def fill_cell(cell, text, *, bold=False, italic=False, color=INK, fill=None, align="left", size=10):
    cell.text = ""
    set_cell_margins(cell)
    if fill:
        shade_cell(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    style_paragraph(p, space_before=2, space_after=2, line=1.12)
    add_text(p, text, size=size, bold=bold, italic=italic, color=color)


def apply_brand_chrome(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run()
    run.add_picture(str(ASSETS / "header.png"), width=Inches(7.15))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp, "unnaturalintelligence.club", size=8, color=SLATE, name="JetBrains Mono")
    mp = footer.add_paragraph()
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mrun = mp.add_run()
    mrun.add_picture(str(ASSETS / "jellyfish-mini.png"), width=Inches(0.38))


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    for i in range(1, 4):
        h = doc.styles[f"Heading {i}"]
        h.font.name = FONT
        h.font.bold = True
        h.font.color.rgb = INK


def build() -> Path:
    if not (ASSETS / "header.png").exists():
        export_kit()

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(28)
    section.bottom_margin = Mm(24)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(4)
    apply_brand_chrome(section)

    core = doc.core_properties
    core.title = "Unnatural Intelligence Event Proposal Template"
    core.author = "Unnatural Intelligence"
    core.subject = "Saitsa student club event proposal — Fall 2026"

    picture(doc, "jellyfish.png", 2.1)
    picture(doc, "wordmark.png", 3.4)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, space_after=4)
    add_text(title, "Event Proposal", size=28, bold=True)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(date, "Template  ·  Fall 2026", size=13, bold=True, color=WINE, name="JetBrains Mono")
    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(tag, copy.PILLARS, size=11, italic=True, color=SLATE)
    picture(doc, "tagline.png", 3.2)
    body(
        doc,
        "Saitsa Clubs Portal  ·  portal.hellorubric.com",
        italic=False,
        color=SLATE,
        size=10,
        center=True,
    )
    signature(doc)

    doc.add_page_break()
    heading(doc, "[EVENT NAME] Overview", center=True)
    heading(doc, "Description", level=2)
    body(doc, "[Brief description of the event here]", italic=True, color=SLATE)
    heading(doc, "Time and Date", level=2)
    body(doc, "[DAY, DATE  ·  START–END]", italic=True, color=SLATE)
    heading(doc, "Proposed Location", level=2)
    body(
        doc,
        "Preferred room? Indoor/Outdoor? Offsite? SECC books campus rooms.",
        italic=True,
        color=SLATE,
    )
    heading(doc, "Expected Attendance", level=2)
    body(doc, "5–10?  20?  50?  100?  10000000????", italic=True, color=SLATE)
    signature(doc)

    doc.add_page_break()
    heading(doc, "Agenda", center=True)
    body(
        doc,
        "Include the important segments. Estimate if you cannot give exact times.",
        italic=True,
        color=SLATE,
        center=True,
    )
    agenda = doc.add_table(rows=7, cols=2)
    agenda.autofit = False
    set_table_borders(agenda)
    col_w = [2200, 7300]
    headers = ["Duration", "Segment"]
    agenda_rows = [
        ["[DURATION]", "[SEGMENT — e.g. welcome, club presentation, member sign-ups]"],
        ["[DURATION]", "[SEGMENT — topic, format]"],
        ["[DURATION]", "[SEGMENT — topic, format]"],
        ["[DURATION]", "[SEGMENT — topic, format]"],
        ["[DURATION]", "[SEGMENT — e.g. Q&A and open discussion]"],
        ["[DURATION]", "[SEGMENT]"],
    ]
    for i, header in enumerate(headers):
        cell = agenda.rows[0].cells[i]
        set_cell_width(cell, col_w[i])
        fill_cell(cell, header, bold=True, color=WHITE, fill=INK_HEX)
    for r, row in enumerate(agenda_rows, start=1):
        for i, text in enumerate(row):
            cell = agenda.rows[r].cells[i]
            set_cell_width(cell, col_w[i])
            fill_cell(cell, text, italic=True, color=SLATE, fill=PAPER_HEX)
    signature(doc)

    doc.add_page_break()
    heading(doc, "Objectives", center=True)
    body(
        doc,
        "What this event is for. Tie at least one objective to a club foundation.",
        italic=True,
        color=SLATE,
        center=True,
    )
    for prompt in [
        "1. [OBJECTIVE 1 — e.g. what students will learn or experience]",
        "2. [OBJECTIVE 2 — e.g. hands-on practice for every attendee]",
        "3. [OBJECTIVE 3 — e.g. recruit new club members]",
        "4. [OBJECTIVE 4 — e.g. promote a foundation (untapped / workflow / human finish) and feed the pipeline of future events]",
    ]:
        body(doc, prompt, italic=True, color=SLATE, after=16)
    signature(doc)

    doc.add_page_break()
    heading(doc, "Requested Supplies and Budgeting", center=True)
    body(
        doc,
        "List every supply, including promotional materials. Built-in costs of "
        f"${PRINTING_COST:.0f} printing ({POSTER_QTY} posters + {NAMETAG_QTY} name tags) "
        f"and ${MISC_COST:.0f} misc are already on the table — add anything else "
        "this event needs and update the total.",
    )
    body(doc, "Attach quotes for all purchases.", bold=True)
    budget = doc.add_table(rows=7, cols=3)
    budget.autofit = False
    set_table_borders(budget)
    b_w = [2100, 5400, 2000]
    b_headers = ["Item", "Details", "Estimated cost"]
    b_rows = [
        ("Printing expenses", PRINTING_DETAILS, money(PRINTING_COST), LIME_HEX, False),
        ("Misc", "Tape, markers, unexpected small costs", money(MISC_COST), LIME_HEX, False),
        ("[ITEM]", "[DETAILS]", "[COST]", PAPER_HEX, True),
        ("[ITEM]", "[DETAILS]", "[COST]", PAPER_HEX, True),
        ("[ITEM]", "[DETAILS]", "[COST]", PAPER_HEX, True),
        (
            "TOTAL",
            "Built-in floor. Add event-specific items above and update this total.",
            money(BUILT_IN_TOTAL),
            PINK_HEX,
            False,
        ),
    ]
    for i, header in enumerate(b_headers):
        cell = budget.rows[0].cells[i]
        set_cell_width(cell, b_w[i])
        fill_cell(
            cell,
            header,
            bold=True,
            color=WHITE,
            fill=INK_HEX,
            align="right" if i == 2 else "left",
        )
    for r, (item, details, cost, fill, placeholder) in enumerate(b_rows, start=1):
        values = (item, details, cost)
        for i, text in enumerate(values):
            cell = budget.rows[r].cells[i]
            set_cell_width(cell, b_w[i])
            fill_cell(
                cell,
                text,
                bold=(i == 0 and not placeholder) or item == "TOTAL",
                italic=placeholder,
                color=WINE if item == "TOTAL" else (SLATE if placeholder else INK),
                fill=fill,
                align="right" if i == 2 else "left",
            )
    bullet(doc, f"Lime rows are built into every proposal (${PRINTING_COST:.0f} printing + ${MISC_COST:.0f} misc = ${BUILT_IN_TOTAL:.0f}). Do not delete them.")
    bullet(doc, "Funding: Saitsa Clubs Grant or start-up reimbursement. Do not spend before approval.")
    bullet(doc, "Eligible spend only — no alcohol, no gifts for members, swag capped where the handbook says.")
    signature(doc)

    doc.add_page_break()
    heading(doc, "Logistics, Risk & Compliance", center=True)
    for text in [
        "Room booking coordinated exclusively by the SECC — no direct booking by the club.",
        "[RISK LEVEL — e.g. low-risk: on campus, indoor, no alcohol, no physical activity]",
        "If refreshments are served: food and beverages ordered through SAIT’s approved catering vendor only, arranged via the SECC.",
        "A/V equipment (if any) picked up from SECC and returned [WHEN — e.g. the same evening].",
        "Primary Event Organizer (PEO) will be present for the full duration of the event and will manage risk.",
        "Attendance list will be kept for the event report.",
        f"Printing expenses (${PRINTING_COST:.0f}): {POSTER_QTY} tabloid 11×17 posters (${POSTER_COST:.0f}) and {NAMETAG_QTY} name tags (${NAMETAG_COST:.0f}). Posters: Saitsa Clubs logo unaltered, submitted via saitsa.com/Poster-Submission before posting.",
        "Posters removed by the club within 48 hours after the event.",
    ]:
        bullet(doc, text)
    heading(doc, "Promotion Plan", center=True)
    for text in [
        "Posters on Saitsa General Poster Boards (approved submission first).",
        "Announcements in club channels and student groups.",
        "Promotion at [CLUBS EVENT — e.g. Clubs Fall Showcase] if the event is scheduled after it.",
        "Word of mouth in classes; open to all SAIT students to maximize reach.",
    ]:
        bullet(doc, text)
    picture(doc, "saitsa-clubs.png", 1.3)
    signature(doc)

    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
