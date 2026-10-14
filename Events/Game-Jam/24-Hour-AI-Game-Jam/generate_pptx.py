#!/usr/bin/env python3
"""A4 PowerPoint of the 24-hour AI Game Jam proposal (Google Slides / PowerPoint)."""

from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
EVENTS = HERE.parent
REPO = EVENTS.parent
CHARTER = REPO / "charter"
QUOTES = HERE / "quotes"
sys.path.insert(0, str(CHARTER))

from revamp_kit import ASSETS, export as export_kit  # noqa: E402

OUTPUT = HERE / "event-proposal-24-hour-AI-game-jam.pptx"

INK = RGBColor(0x03, 0x02, 0x05)
WINE = RGBColor(0x59, 0x13, 0x37)
SLATE = RGBColor(0x5A, 0x5A, 0x5A)
PAPER = RGBColor(0xF9, 0xF9, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PINK = RGBColor(0xFA, 0x4A, 0x9E)
LIME = RGBColor(0xBA, 0xF8, 0x37)
RULE = RGBColor(0xD7, 0xD7, 0xDA)
LIME_WASH = RGBColor(0xE8, 0xF6, 0xC4)
PINK_WASH = RGBColor(0xFA, 0xD6, 0xE8)
INK_HEX = "030205"
RULE_HEX = "D7D7DA"

PAGE_W = Inches(8.27)
PAGE_H = Inches(11.69)
M = Inches(0.55)
W = Inches(7.17)
CONTENT_TOP = Inches(1.12)
FOOTER_Y = Inches(11.18)

POSTER = 10.00
MISC = 5.00


def money(value: float) -> str:
    return f"${value:,.2f}"


def _prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = PAGE_W
    prs.slide_height = PAGE_H
    prs.core_properties.title = "24-Hour AI Game Jam — Event Proposal"
    prs.core_properties.author = "Unnatural Intelligence"
    prs.core_properties.subject = "Saitsa student club event proposal — Fall 2026"
    return prs


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def fill_slide(slide, color=PAPER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), PAGE_W, PAGE_H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def chrome(slide, prs, page: int) -> None:
    fill_slide(slide)
    slide.shapes.add_picture(
        str(ASSETS / "header.png"),
        Inches(0.16),
        Inches(0.08),
        width=prs.slide_width - Inches(0.32),
    )
    box = slide.shapes.add_textbox(Inches(0), FOOTER_Y, prs.slide_width, Inches(0.22))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"{page:02d}"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = WINE
    run.font.name = "JetBrains Mono"
    slide.shapes.add_picture(
        str(ASSETS / "jellyfish-mini.png"),
        prs.slide_width / 2 - Inches(0.18),
        Inches(11.28),
        width=Inches(0.36),
    )


def tb(
    slide,
    l,
    t,
    w,
    h,
    text,
    *,
    size=12,
    bold=False,
    italic=False,
    color=INK,
    align=PP_ALIGN.LEFT,
    font="Inter",
):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    parts = str(text).split("\n")
    for i, part in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = part
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    return tf


def title(slide, text, *, size=20):
    tb(slide, M, CONTENT_TOP, W, Inches(0.42), text, size=size, bold=True, align=PP_ALIGN.CENTER)


def signature_line(slide, named=False):
    y = Inches(10.52)
    peo = (
        "Primary Event Organizer — VP, Unnatural Intelligence (Pierre Tchio)"
        if named
        else "Primary Event Organizer — [NAME AND EXECUTIVE ROLE]"
    )
    tb(slide, M, y, W, Inches(0.18), peo, size=8, bold=True, align=PP_ALIGN.CENTER)
    tb(
        slide,
        M,
        y + Inches(0.18),
        W,
        Inches(0.18),
        "Signature: ____________________     Date: ______________",
        size=8,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )


def rounded(slide, l, t, w, h, fill=WHITE, line=RULE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.75)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def card(slide, l, t, w, h, heading, body, *, body_size=11):
    rounded(slide, l, t, w, h)
    pad = Inches(0.14)
    body_top = t + Inches(0.10)
    body_h = h - Inches(0.18)
    if heading:
        tb(slide, l + pad, t + Inches(0.08), w - pad * 2, Inches(0.24), heading, size=12, bold=True)
        body_top = t + Inches(0.32)
        body_h = h - Inches(0.40)
    tb(slide, l + pad, body_top, w - pad * 2, body_h, body, size=body_size, color=INK)


def _cell_fill(cell, color: RGBColor) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = color


def _cell_border(cell, color_hex=RULE_HEX, width_pt=0.5) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    ln_w = str(int(width_pt * 12700))
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        for child in list(tcPr):
            if child.tag == qn(f"a:{edge}"):
                tcPr.remove(child)
        ln = etree.SubElement(tcPr, qn(f"a:{edge}"))
        ln.set("w", ln_w)
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        solid = etree.SubElement(ln, qn("a:solidFill"))
        srgb = etree.SubElement(solid, qn("a:srgbClr"))
        srgb.set("val", color_hex)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")


def _set_cell(
    cell,
    text,
    *,
    size=9,
    bold=False,
    color=INK,
    fill=WHITE,
    align=PP_ALIGN.LEFT,
    font="Inter",
):
    _cell_fill(cell, fill)
    _cell_border(cell)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    p.clear()
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_table(
    slide,
    left,
    top,
    width,
    headers,
    rows,
    col_fracs,
    *,
    aligns=None,
    row_h=0.30,
    header_h=0.28,
    font_size=9,
):
    """rows: list of (cells, kind) where kind is body | builtin | total."""
    aligns = aligns or ["left"] * len(headers)
    n_cols = len(headers)
    n_rows = 1 + len(rows)
    height = Inches(header_h + row_h * len(rows))
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    for i, frac in enumerate(col_fracs):
        table.columns[i].width = Emu(int(width * frac))
    table.rows[0].height = Inches(header_h)
    for j, header in enumerate(headers):
        _set_cell(
            table.cell(0, j),
            header,
            size=8.5,
            bold=True,
            color=PAPER,
            fill=INK,
            align=PP_ALIGN.RIGHT if aligns[j] == "right" else PP_ALIGN.LEFT,
            font="Inter",
        )
        _cell_border(table.cell(0, j), INK_HEX, 0.5)
    fills = {"body": WHITE, "builtin": LIME_WASH, "total": PINK_WASH, "blank": WHITE}
    for r, (cells, kind) in enumerate(rows, start=1):
        table.rows[r].height = Inches(row_h)
        fill = fills[kind]
        text_color = WINE if kind == "total" else INK
        for c, text in enumerate(cells):
            is_emph = kind in {"builtin", "total"} and c == 0 or kind == "total"
            _set_cell(
                table.cell(r, c),
                text,
                size=font_size,
                bold=is_emph,
                color=text_color,
                fill=fill,
                align=PP_ALIGN.RIGHT if aligns[c] == "right" else PP_ALIGN.LEFT,
            )
    return Inches(top.inches + header_h + row_h * len(rows))


def add_picture_fit(slide, path: Path, left, top, max_w, max_h):
    with Image.open(path) as im:
        iw, ih = im.size
    max_w_in = max_w.inches
    max_h_in = max_h.inches
    aspect = ih / iw
    w_in = max_w_in
    h_in = w_in * aspect
    if h_in > max_h_in:
        h_in = max_h_in
        w_in = h_in / aspect
    x_in = left.inches + (max_w_in - w_in) / 2
    frame = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x_in),
        top,
        Inches(w_in),
        Inches(h_in),
    )
    frame.fill.background()
    frame.line.color.rgb = RULE
    frame.line.width = Pt(0.5)
    slide.shapes.add_picture(str(path), Inches(x_in), top, width=Inches(w_in), height=Inches(h_in))
    return Inches(h_in)


def quote_block(slide, filename, left, top, max_w, max_h, caption):
    h = add_picture_fit(slide, QUOTES / filename, left, top, max_w, max_h)
    cap_top = Inches(top.inches + h.inches + 0.05)
    tb(slide, left, cap_top, max_w, Inches(0.38), caption, size=8.5, italic=True, color=SLATE)
    return Inches(cap_top.inches + 0.40)


def badge_row(slide, slugs, y):
    paths = [ASSETS / f"badge-{slug}.png" for slug in slugs]
    widths = []
    for path in paths:
        with Image.open(path) as im:
            widths.append(im.width / im.height * 0.28)
    gap = 0.10
    total = sum(widths) + gap * (len(widths) - 1)
    x = (PAGE_W.inches - total) / 2
    for path, w in zip(paths, widths):
        slide.shapes.add_picture(str(path), Inches(x), y, height=Inches(0.28))
        x += w + gap


def budget_rows(items, total_note):
    rows = []
    total = 0.0
    for item, details, cost, kind in items:
        total += cost
        rows.append(([item, details, money(cost)], kind))
    rows.append((["TOTAL", total_note, money(total)], "total"))
    return rows


def page_cover(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    s.shapes.add_picture(str(ASSETS / "jellyfish.png"), Inches(3.05), Inches(1.35), height=Inches(2.4))
    s.shapes.add_picture(str(ASSETS / "wordmark.png"), Inches(1.9), Inches(3.85), width=Inches(4.5))
    tb(s, Inches(0.4), Inches(5.45), Inches(7.5), Inches(0.5), "Event Proposal", size=28, bold=True, align=PP_ALIGN.CENTER)
    tb(
        s,
        Inches(0.4),
        Inches(6.00),
        Inches(7.5),
        Inches(0.32),
        "24-Hour AI Game Jam  ·  Fall 2026",
        size=13,
        bold=True,
        color=WINE,
        align=PP_ALIGN.CENTER,
        font="JetBrains Mono",
    )
    tb(
        s,
        Inches(0.4),
        Inches(6.38),
        Inches(7.5),
        Inches(0.30),
        "Unnatural Intelligence  ×  Game Development Club",
        size=12,
        italic=True,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )
    tb(
        s,
        Inches(0.4),
        Inches(6.72),
        Inches(7.5),
        Inches(0.28),
        "Godot 4  ·  paid Ziva Basic  ·  24 hours  ·  doors Sat 9:00",
        size=11,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )
    s.shapes.add_picture(str(ASSETS / "tagline.png"), Inches(2.4), Inches(7.15), width=Inches(3.5))
    badge_row(s, ["workshop", "live-lab", "halt-restart", "clubs-expo"], Inches(8.05))
    tb(
        s,
        Inches(0.4),
        Inches(8.50),
        Inches(7.5),
        Inches(0.22),
        "Saitsa Clubs Portal  ·  portal.hellorubric.com",
        size=9,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )
    signature_line(s, named=True)


def page_overview(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "24-Hour AI Game Jam Overview")
    y = 1.62
    blocks = [
        (
            "Description",
            "Overnight Saturday–Sunday jam co-hosted with the Game Development Club. "
            "Doors Saturday 9:00. Theme, teams, Godot 4, and Ziva Basic before 10:30. "
            "Official build is 24 hours (Sat 10:30 – Sun 10:30). Teams of 2–5 ship a "
            "playable Godot prototype. Food is five drops from the attached Curated "
            "Catering campus menu only — no Event Palette items. Not a sleepover: "
            "leave-and-return allowed; a named exec is on site all night. No alcohol.",
            1.72,
        ),
        (
            "Time and Date",
            "Saturday 14 November – Sunday 15 November 2026. Doors Sat 9:00 AM. "
            "Official jam Sat 10:30 AM – Sun 10:30 AM. Judging then teardown to ~1:00 PM Sunday. "
            "Pending SECC, Campus Security, and the 5 October grant intake.",
            1.05,
        ),
        (
            "Proposed Location",
            "Indoor computer-friendly classroom or lab with projector, power, and Wi-Fi "
            "(Stan Grad / Heritage Hall). Weekend and after-hours booking by SECC only. "
            "Overflow if RSVPs exceed ~40.",
            0.95,
        ),
        (
            "Expected Attendance",
            "30 recommended (fully fed, Philippines dinner) / 50 ceiling (pizza dinner). "
            "Engine: Godot 4 (free). AI: Ziva Basic at CAD $28/attendee. Cursor is BYO only. "
            "PEO: VP, Unnatural Intelligence (Pierre Tchio). Game Dev Club is logistics lead "
            "and overnight deputy. Equal budget split.",
            1.18,
        ),
    ]
    for heading, body, h in blocks:
        tb(s, M, Inches(y), W, Inches(0.24), heading, size=12, bold=True)
        card(s, M, Inches(y + 0.24), W, Inches(h), "", body, body_size=11)
        # heading is outside the card like the PDF label; body sits in the card.
        # Re-draw body only — empty heading inside card wastes space. Use inner pad.
        y += 0.24 + h + 0.10
    signature_line(s)


def page_agenda(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Agenda — Saturday 9:00 to Sunday afternoon", size=18)
    tb(
        s,
        M,
        Inches(1.55),
        W,
        Inches(0.40),
        "24 hours of game dev is fixed. Official clock Sat 10:30 – Sun 10:30. Quiet hours 01:00–06:00.",
        size=11,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )
    add_table(
        s,
        M,
        Inches(2.05),
        W,
        ["Clock", "Segment"],
        [
            (["Sat 09:00–09:15", "Doors, attendance, power map, both-club sign-ups"], "body"),
            (["Sat 09:15–10:00", "Welcome, rules, secret theme, teams (2–5)"], "body"),
            (["Sat 09:30", "Continental Lite kickoff ($10) — coffee included"], "body"),
            (["Sat 10:00–10:30", "Setup lab: Godot 4 + Ziva Basic redeem"], "body"),
            (["Sat 10:30", "Official start. 24-hour clock begins."], "builtin"),
            (["Sat 12:30", "Lunch boxes drop ($16). Eat at desks."], "body"),
            (["Sat 18:30", "Philippines dinner ($29) · pizza ($6) if 50-person ceiling"], "body"),
            (["Sun 00:00", "Energy bars & bites ($6) midnight drop"], "body"),
            (["Sun 01:00–06:00", "Quiet hours. Leave-and-return. Two-person exec minimum"], "body"),
            (["Sun 08:00", "Signature Breakfast Lite ($12)"], "body"),
            (["Sun 10:30", "Submissions freeze. Godot export + 60-second playable"], "builtin"),
            (["Sun 10:30–13:00", "Judging, results, send-home, exec teardown"], "body"),
        ],
        [0.28, 0.72],
        row_h=0.52,
        header_h=0.32,
        font_size=11,
    )
    signature_line(s)


def page_objectives(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Objectives & labour split")
    prompts = [
        "Ship a playable Godot game in 24 hours — every team leaves with a running scene.",
        "Use paid Ziva Basic as a Godot teammate, not as the designer — human finish.",
        "Cross-club recruitment — both clubs sign people up; beginners pair with shippers.",
        "Foundations: untapped (AI in Godot overnight), workflow (a real jam), human finish.",
    ]
    y = 1.62
    for i, prompt in enumerate(prompts, start=1):
        rounded(s, M, Inches(y), W, Inches(0.72))
        circle = s.shapes.add_shape(
            MSO_SHAPE.OVAL,
            M + Inches(0.14),
            Inches(y + 0.18),
            Inches(0.36),
            Inches(0.36),
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = LIME if i % 2 else PINK
        circle.line.fill.background()
        tb(
            s,
            M + Inches(0.14),
            Inches(y + 0.22),
            Inches(0.36),
            Inches(0.30),
            str(i),
            size=12,
            bold=True,
            align=PP_ALIGN.CENTER,
            font="JetBrains Mono",
        )
        tb(s, M + Inches(0.62), Inches(y + 0.16), W - Inches(0.76), Inches(0.46), prompt, size=12)
        y += 0.82
    tb(s, M, Inches(y + 0.04), W, Inches(0.32), "Who does what", size=14, bold=True, align=PP_ALIGN.CENTER)
    add_table(
        s,
        M,
        Inches(y + 0.38),
        W,
        ["Unnatural Intelligence", "Game Development Club"],
        [
            (["Files proposal + grant; 50% of cost", "Coordinates logistics via SECC; 50% of cost"], "body"),
            (["Ziva Basic redeem + setup lab", "Godot install, template, export"], "body"),
            (["First draft of judging rubric", "Reviews rubric; co-judges"], "body"),
            (["Named exec on overnight roster", "Named exec on overnight roster"], "body"),
            (["Theme together (secret until Sat 9:15)", "Theme together"], "builtin"),
        ],
        [0.5, 0.5],
        row_h=0.48,
        header_h=0.32,
        font_size=11,
    )
    signature_line(s)


def page_menu(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Campus menu only — attached cards", size=18)
    tb(
        s,
        M,
        Inches(1.52),
        W,
        Inches(0.78),
        "Food quotes are the attached Curated Catering cards, not the Hotel Arts Event Palette. "
        "Order via SECC. SAITevents@hotelarts.ca · 403-210-5774. "
        "Fees card: +18% service + 5% GST (× 1.239). 20% / 50% deposits after grant approval. "
        "Additional staff $30/hr not budgeted. No statutory-holiday surcharge on 14–15 Nov 2026.",
        size=11,
        color=SLATE,
    )
    y = add_table(
        s,
        M,
        Inches(2.32),
        W,
        ["Menu card", "Base", "Loaded", "This jam"],
        [
            (["Continental Lite", "$10.00", "$12.39", "Sat 9:30 kickoff"], "body"),
            (["Lunch boxes", "$16.00", "$19.82", "Sat 12:30 lunch"], "builtin"),
            (["Taste of the Philippines", "$29.00", "$35.93", "Sat dinner (30 rec.)"], "builtin"),
            (["Taste of Morocco", "$30.00", "$37.17", "Dinner swap (~+$37 / 30)"], "body"),
            (["Pizza, 2 slices", "$6.00", "$7.43", "Sat dinner (50 ceiling)"], "body"),
            (["Energy bars & bites", "$6.00", "$7.43", "Sun 00:00 midnight"], "body"),
            (["Signature Breakfast Lite", "$12.00", "$14.87", "Sun 8:00 breakfast"], "body"),
            (["Sides & add-ons", "$1–$6", "—", "Not used"], "body"),
        ],
        [0.38, 0.14, 0.16, 0.32],
        aligns=["left", "right", "right", "left"],
        row_h=0.38,
        header_h=0.30,
        font_size=10,
    )
    tb(
        s,
        M,
        Inches(y.inches + 0.16),
        W,
        Inches(0.28),
        "Paid Ziva Basic — Cursor is not a grant line",
        size=14,
        bold=True,
    )
    tb(
        s,
        M,
        Inches(y.inches + 0.44),
        W,
        Inches(0.95),
        "Hobby ($3 usage) will not last 24 hours. Basic is $20 USD/mo = CAD $28.00 "
        "(BoC USD/CAD 1.4002 on 18 Sep 2026) with $20 weekly / $60 monthly usage. "
        "30 × $28 = $840. 50 × $28 = $1,400. Cancel after the jam. "
        "Stay under the weekly cap with Ziva’s free model plus any attendee’s own Claude Code / Codex. "
        "Cursor Pro is the same $20 USD — do not buy 30–50 seats.",
        size=11,
        color=SLATE,
    )
    add_table(
        s,
        M,
        Inches(y.inches + 1.42),
        W,
        ["Headcount", "Ziva Basic @ $28", "Five food drops loaded"],
        [
            (["30", "$840", "$2,714 (Philippines dinner)"], "builtin"),
            (["50", "$1,400", "$3,098 (pizza dinner)"], "body"),
        ],
        [0.28, 0.32, 0.40],
        aligns=["left", "right", "right"],
        row_h=0.40,
        header_h=0.30,
        font_size=11,
    )
    signature_line(s)


def page_budget_30(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Budget A — $4,100 for 30 people", size=18)
    tb(
        s,
        M,
        Inches(1.55),
        W,
        Inches(0.52),
        "Recommended. Five campus-menu drops including Philippines dinner, Ziva Basic for 30, "
        "overnight security placeholder. Split $2,050 / $2,050.",
        size=11,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )
    add_table(
        s,
        M,
        Inches(2.12),
        W,
        ["Item", "Details", "Estimated cost"],
        budget_rows(
            [
                ("Posters", "Tabloid 11×17, unaltered Saitsa Clubs logo", POSTER, "builtin"),
                ("Misc", "Tape, markers, unexpected small costs", MISC, "builtin"),
                ("Continental Lite (30)", "Sat kickoff @ $10 pp + 18% + GST", 372, "body"),
                ("Lunch boxes (30)", "Sat lunch @ $16 pp loaded", 595, "body"),
                ("Philippines (30)", "Sat dinner @ $29 pp loaded", 1078, "body"),
                ("Energy bars (30)", "Midnight @ $6 pp loaded", 223, "body"),
                ("Signature Breakfast (30)", "Sun @ $12 pp loaded", 446, "body"),
                ("Ziva Basic (30)", "CAD $28 / attendee, one month then cancel", 840, "body"),
                ("After-hours security", "Placeholder until Security quote", 350, "body"),
                ("Name tags", "Printed club templates", 30, "body"),
                ("Power strips", "Overnight laptop load", 50, "body"),
                ("First aid + buffer", "Overnight kit + dietary / minimums", 101, "body"),
            ],
            "UI $2,050 + Game Dev $2,050. Do not spend before approval.",
        ),
        [0.30, 0.50, 0.20],
        aligns=["left", "left", "right"],
        row_h=0.48,
        header_h=0.32,
        font_size=10,
    )
    signature_line(s)


def page_budget_50(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Budget B — $5,000 for 50 people", size=18)
    tb(
        s,
        M,
        Inches(1.55),
        W,
        Inches(0.58),
        "Ceiling. Same five drops but pizza dinner ($6) instead of Philippines ($29) — "
        "the $29 buffet at 50 people is $1,797 loaded and does not fit with Ziva Basic. "
        "Split $2,500 / $2,500. Do not go above 50.",
        size=11,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )
    add_table(
        s,
        M,
        Inches(2.18),
        W,
        ["Item", "Details", "Estimated cost"],
        budget_rows(
            [
                ("Posters", "Tabloid 11×17, unaltered Saitsa Clubs logo", POSTER, "builtin"),
                ("Misc", "Tape, markers, unexpected small costs", MISC, "builtin"),
                ("Continental Lite (50)", "Sat kickoff @ $10 pp loaded", 620, "body"),
                ("Lunch boxes (50)", "Sat lunch @ $16 pp loaded", 991, "body"),
                ("Pizza (50)", "Sat dinner @ $6 pp loaded (2 slices)", 372, "body"),
                ("Energy bars (50)", "Midnight @ $6 pp loaded", 372, "body"),
                ("Signature Breakfast (50)", "Sun @ $12 pp loaded", 743, "body"),
                ("Ziva Basic (50)", "CAD $28 / attendee, one month then cancel", 1400, "body"),
                ("After-hours security", "Placeholder until Security quote", 250, "body"),
                ("Name tags", "Printed club templates", 50, "body"),
                ("Power strips", "Higher laptop load", 60, "body"),
                ("First aid + buffer", "Overnight kit + dietary / minimums", 127, "body"),
            ],
            "UI $2,500 + Game Dev $2,500. Do not spend before approval.",
        ),
        [0.30, 0.50, 0.20],
        aligns=["left", "left", "right"],
        row_h=0.48,
        header_h=0.32,
        font_size=10,
    )
    signature_line(s)


def page_quotes_food_1(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Quoted menu — breakfast & lunch", size=18)
    tb(
        s,
        M,
        Inches(1.52),
        W,
        Inches(0.36),
        "Attached Curated Catering cards. Base prices; add 18% + GST (× 1.239).",
        size=11,
        color=SLATE,
        align=PP_ALIGN.CENTER,
    )
    y = quote_block(
        s,
        "buffet-breakfast.png",
        M,
        Inches(1.92),
        W,
        Inches(3.55),
        "Continental Lite $10 (Sat kickoff) · Signature Breakfast Lite $12 (Sun 8:00). Min 15. Coffee included.",
    )
    quote_block(
        s,
        "lunch-boxes.png",
        M,
        y,
        W,
        Inches(3.55),
        "Lunch boxes $16. Min 5. 21–50 guests = 2 protein choices. Sat 12:30.",
    )
    signature_line(s)


def page_quotes_food_2(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Quoted menu — dinner & midnight", size=18)
    y = quote_block(
        s,
        "taste-philippines-morocco.png",
        M,
        Inches(1.58),
        W,
        Inches(2.15),
        "Philippines $29 (30-person dinner) · Morocco $30 (swap). Min 15 guests.",
    )
    y = quote_block(
        s,
        "pizza.png",
        M,
        y,
        W,
        Inches(1.95),
        "Pizza $6 / 2 slices. 50-person ceiling dinner. Max four flavours.",
    )
    y = quote_block(
        s,
        "energy-bars.png",
        M,
        y,
        W,
        Inches(1.45),
        "Energy bars & bites $6 / 3 pieces. Midnight drop. Min 5 guests.",
    )
    quote_block(
        s,
        "sides-addons.png",
        M,
        y,
        W,
        Inches(1.20),
        "Sides & add-ons — not used in these budgets.",
    )
    signature_line(s)


def page_quotes_fees_ai(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Quoted fees, Ziva, FX", size=18)
    y = quote_block(
        s,
        "deposit-fees.png",
        M,
        Inches(1.55),
        W,
        Inches(3.15),
        "18% service + 5% GST on all food. 20% / 50% deposits after grant approval. Staff $30/hr not budgeted.",
    )
    y = quote_block(
        s,
        "ziva-usage-limits.png",
        M,
        y,
        W,
        Inches(2.85),
        "Ziva Basic $20 USD / $60 usage — the paid SKU for every attendee. Hobby $3 is too small for 24 hours.",
    )
    quote_block(
        s,
        "boc-usd-cad.png",
        M,
        y,
        W,
        Inches(1.55),
        "BoC FXUSDCAD 18 Sep 2026 = 1.4002 → USD $20 = CAD $28.00.",
    )
    signature_line(s)


def page_logistics(prs, page):
    s = blank(prs)
    chrome(s, prs, page)
    title(s, "Logistics, Risk & Compliance", size=18)
    def bullet_list(top, height, lines):
        box = s.shapes.add_textbox(M, Inches(top), W, Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        for i, text in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"•  {text}"
            p.level = 0
            p.font.size = Pt(12)
            p.font.color.rgb = INK
            p.font.name = "Inter"
            p.space_after = Pt(7)

    bullet_list(
        1.55,
        4.05,
        [
            "Room booking coordinated exclusively by the SECC. Game Dev logistics lead files through them.",
            "Risk: medium. Indoor, no alcohol, no sport — after-hours, fatigue, food, laptops. Written exec roster; PEO or deputy always on site.",
            "After-hours with SECC and Campus Security. Occupants staying after facility hours listed. Booked work session, not an unsanctioned sleepover. Quiet hours 01:00–06:00.",
            "Food only through Curated Catering via SECC, using the attached campus menu cards. Dietary needs on RSVP. Drops 9:30, 12:30, 18:30, 00:00, 08:00.",
            "A/V returned after Sunday teardown (~13:00). Attendance list signed. Posters tabloid 11×17, Saitsa Clubs logo unaltered, $10 built-in.",
            "Disclose the 50/50 sister-club contribution. No spend before approval. Unused grant (including unspent security) is reclaimed.",
            "Ziva Basic is one month of event access, then cancelled — not a member gift. Do not buy Cursor for the room.",
        ],
    )
    tb(s, M, Inches(5.72), W, Inches(0.32), "Promotion Plan", size=16, bold=True, align=PP_ALIGN.CENTER)
    bullet_list(
        6.08,
        2.15,
        [
            "Posters on Saitsa General Poster Boards (approved submission first).",
            "Both clubs’ Discord, Instagram, Teams, and Rubric RSVP. Cap 50; waitlist after that.",
            "Clubs Fall Expo / Showcase (13–14 October 2026) if the jam is after it.",
            "Pitch: Godot + Ziva Basic, 24 hours, doors Saturday 9:00, meals included, beginners welcome.",
        ],
    )
    s.shapes.add_picture(
        str(ASSETS / "saitsa-clubs.png"),
        PAGE_W - M - Inches(1.15),
        Inches(8.95),
        width=Inches(1.15),
    )
    signature_line(s, named=True)


def build() -> Path:
    if not (ASSETS / "header.png").exists():
        export_kit()

    prs = _prs()
    pages = [
        page_cover,
        page_overview,
        page_agenda,
        page_objectives,
        page_menu,
        page_budget_30,
        page_budget_50,
        page_quotes_food_1,
        page_quotes_food_2,
        page_quotes_fees_ai,
        page_logistics,
    ]
    for i, fn in enumerate(pages, start=1):
        fn(prs, i)
    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
