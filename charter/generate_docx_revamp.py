#!/usr/bin/env python3
"""Editable Word copy of the 2026–27 branded charter."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Mm

import revamp_copy as copy
from revamp_kit import ASSETS, export as export_kit

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Unnatural_Intelligence_Club_Charter_2026-27.docx"
FONT = "Inter"

INK = RGBColor(0x03, 0x02, 0x05)
WINE = RGBColor(0x59, 0x13, 0x37)
SLATE = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SWATCHES = [
    ("#F9F9FA", RGBColor(0xF9, 0xF9, 0xFA), False),
    ("#95B747", RGBColor(0x95, 0xB7, 0x47), False),
    ("#030205", RGBColor(0x03, 0x02, 0x05), True),
    ("#FA4A9E", RGBColor(0xFA, 0x4A, 0x9E), True),
    ("#591337", RGBColor(0x59, 0x13, 0x37), True),
    ("#BAF837", RGBColor(0xBA, 0xF8, 0x37), False),
    ("#7EC8FF", RGBColor(0x7E, 0xC8, 0xFF), False),
]


def set_cell_margins(cell, *, twips=140) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_cell_width(cell, twips: int) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(twips))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, *, size=11, bold=False, italic=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def style_paragraph(p, *, space_before=0, space_after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    return run


def heading(doc, text, level=1, center=False):
    p = doc.add_heading(text, level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(
            run,
            size=22 if level == 1 else 13 if level == 2 else 12,
            bold=True,
            color=INK,
        )
    style_paragraph(p, space_before=14 if level > 1 else 6, space_after=8)
    return p


def hrule(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), "D7D7DA")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=11, bold=True, color=WINE, name="JetBrains Mono")
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run._r.append(instr)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld2)


def apply_brand_chrome(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run()
    run.add_picture(str(ASSETS / "header.png"), width=Inches(7.15))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)
    mp = footer.add_paragraph()
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mrun = mp.add_run()
    mrun.add_picture(str(ASSETS / "jellyfish-mini.png"), width=Inches(0.42))


def body(doc, text, *, italic=False, color=INK, after=8, rule=False):
    p = doc.add_paragraph()
    style_paragraph(p, space_after=after)
    add_text(p, text, size=11, italic=italic, color=color)
    if rule:
        hrule(p)
    return p


def quote(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=12, space_after=6, line=1.3)
    add_text(p, f"“{text}”", size=13, italic=True, color=WINE)
    return p


def bullet(doc, text, level=0, bold=False):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    style_paragraph(p, space_before=1, space_after=2, line=1.12)
    add_text(p, text, size=11, bold=bold)
    return p


def role(doc, title, purpose, items):
    heading(doc, title, level=3)
    p = doc.add_paragraph()
    style_paragraph(p, space_after=4)
    add_text(p, "Core Purpose: ", size=11, bold=True)
    add_text(p, purpose, size=11)
    p = doc.add_paragraph()
    style_paragraph(p, space_before=4, space_after=4)
    add_text(p, "Responsibilities:", size=11, bold=True)
    for text, level in items:
        bullet(doc, text, level=level)


def picture(doc, name, width, center=True):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=6, space_after=8)
    p.add_run().add_picture(str(ASSETS / name), width=Inches(width))
    return p


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    for i in range(1, 4):
        h = doc.styles[f"Heading {i}"]
        h.font.name = FONT
        h.font.bold = True
        h.font.color.rgb = INK


def build() -> Path:
    if not (ASSETS / "header.png").exists():
        export_kit()

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(28)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(4)
    apply_brand_chrome(section)

    core = doc.core_properties
    core.title = "Unnatural Intelligence Club Charter 2026–27"
    core.author = "Unnatural Intelligence"
    core.subject = copy.MISSION_QUOTE

    picture(doc, "jellyfish.png", 2.1)
    picture(doc, "wordmark.png", 3.4)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, space_after=6)
    add_text(title, "Club Charter", size=28, bold=True)
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(date, copy.YEAR, size=13, bold=True, color=WINE, name="JetBrains Mono")
    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(tag, copy.PILLARS, size=11, italic=True, color=SLATE)
    picture(doc, "tagline.png", 3.2)

    doc.add_page_break()
    heading(doc, "Contents", center=True)
    for name, num, bold, _level in copy.TOC:
        if not bold:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        style_paragraph(p, space_after=6)
        add_text(p, name, size=12, bold=True)
        add_text(p, f"\t{num}", size=12, color=SLATE)

    doc.add_page_break()
    heading(doc, "Mission Statement", center=True)
    body(doc, copy.MISSION_LEAD)
    quote(doc, copy.MISSION_QUOTE)
    body(doc, copy.MISSION_FEELINGS)
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    col_w = 3120
    for i, (icon, title_text, para) in enumerate(copy.FOUNDATIONS):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, col_w)
        set_cell_margins(cell, twips=160)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_before=8, space_after=10)
        p.add_run().add_picture(str(ASSETS / f"icon-{icon}.png"), width=Inches(0.42))
        tp = cell.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(tp, space_before=10, space_after=8, line=1.15)
        title_parts = title_text.split("\n")
        add_text(tp, title_parts[0], size=11, bold=True)
        for extra in title_parts[1:]:
            tp.add_run().add_break()
            add_text(tp, extra, size=11, bold=True)
        bp = cell.add_paragraph()
        bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(bp, space_before=4, space_after=8, line=1.2)
        add_text(bp, para, size=10, color=SLATE)
    picture(doc, "quote-useful.png", 2.8)

    doc.add_page_break()
    heading(doc, "The Basis", center=True)
    for title_text, para in copy.BASIS:
        heading(doc, title_text, level=2)
        body(doc, para, rule=True, after=10)

    doc.add_page_break()
    heading(doc, "Goals", center=True)
    for title_text, para in copy.GOALS_A + copy.GOALS_B:
        heading(doc, title_text, level=2)
        body(doc, para, rule=True, after=10)

    doc.add_page_break()
    heading(doc, "How We Operate", center=True)
    body(doc, "From the Saitsa Clubs Handbook 2026/27", italic=True, color=SLATE)
    for title_text, para in copy.OPERATE:
        heading(doc, title_text, level=2)
        body(doc, para, rule=True, after=10)
    body(doc, copy.SIGNING_NOTE, italic=True, color=SLATE)

    doc.add_page_break()
    spacer = doc.add_paragraph()
    spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(spacer, space_before=160)
    add_text(spacer, "Executive Roles & Responsibilities", size=26, bold=True)

    doc.add_page_break()
    heading(doc, "Directors")
    role(doc, *copy.PRESIDENT)
    role(doc, *copy.VP_OPS)
    role(doc, *copy.VP_FINANCE)
    body(
        doc,
        "Saitsa requires every club to keep three key executives filled: "
        "President, Vice-President Operations, and Vice-President Finance.",
        italic=True,
        color=SLATE,
    )

    doc.add_page_break()
    heading(doc, "Marketing")
    role(doc, *copy.SOCIAL)
    role(doc, *copy.DESIGNER)

    doc.add_page_break()
    heading(doc, "Events")
    role(doc, *copy.EVENTS_COORD)
    role(doc, *copy.OUTREACH)

    doc.add_page_break()
    heading(doc, "Event Concepts")
    heading(doc, "Untapped Potential — find what default chat never gets asked.", level=2)
    for text, level in copy.UNTAPPED_ITEMS:
        bullet(doc, text, level=max(level - 1, 0))
    heading(doc, "Workflow Integration — from demo to a Tuesday you would actually keep.", level=2)
    for text, level in copy.WORKFLOW_ITEMS:
        bullet(doc, text, level=max(level - 1, 0))
    heading(doc, "Token Craft & Human Touch — less spend, more of you in the result.", level=2)
    for text, level in copy.TOKEN_ITEMS:
        bullet(doc, text, level=max(level - 1, 0))
    heading(doc, "Mixer events", level=2)
    for text, level in copy.MIXER_ITEMS:
        bullet(doc, text, level=max(level - 1, 0))

    doc.add_page_break()
    heading(doc, "This Year on Campus")
    for date, event in copy.CALENDAR:
        p = doc.add_paragraph()
        style_paragraph(p, space_after=4)
        add_text(p, f"{date}  ", size=11, bold=True, color=WINE, name="JetBrains Mono")
        add_text(p, event, size=11)
    heading(doc, "How we register", level=2)
    for note in copy.REGISTER:
        bullet(doc, note)

    doc.add_page_break()
    heading(doc, "Unnatural Intelligence Style Guide", center=True)
    picture(doc, "lockup-stacked.png", 2.2)
    picture(doc, "jellyfish.png", 1.8)
    body(doc, "Google-friendly fonts: Inter  |  Inter Display  |  JetBrains Mono")
    for note in copy.STYLE_NOTES:
        bullet(doc, note)
    heading(doc, "Copy-paste colour codes", level=2)
    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    for i, (hex_code, rgb, light_text) in enumerate(SWATCHES):
        cell = table.rows[i // 4].cells[i % 4]
        shade_cell(cell, hex_code)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, hex_code, size=10, bold=True, color=WHITE if light_text else INK, name="JetBrains Mono")
    table.rows[1].cells[3].text = ""

    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
