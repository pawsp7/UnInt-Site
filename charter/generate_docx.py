#!/usr/bin/env python3
"""Generate an editable Word (.docx) copy of the club charter.

Google Docs opens .docx natively: Drive → New → File upload, then
Open with → Google Docs. Headings become Docs heading styles so you can
Insert → Table of contents after import.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Mm, Emu

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "Unnatural_Intelligence_Club_Charter.docx"
FONT = "Inter"

TEAL = RGBColor(0x1A, 0x7A, 0x84)
INK = RGBColor(0x1A, 0x1A, 0x1A)
SLATE = RGBColor(0x5A, 0x5A, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SWATCHES = [
    ("#1A7A84", RGBColor(0x1A, 0x7A, 0x84), True),
    ("#8EE3D8", RGBColor(0x8E, 0xE3, 0xD8), False),
    ("#5A5A5A", RGBColor(0x5A, 0x5A, 0x5A), True),
    ("#7B61FF", RGBColor(0x7B, 0x61, 0xFF), True),
    ("#C4B8FF", RGBColor(0xC4, 0xB8, 0xFF), False),
    ("#EAEAEA", RGBColor(0xEA, 0xEA, 0xEA), False),
    ("#FF6B4A", RGBColor(0xFF, 0x6B, 0x4A), True),
    ("#F4C430", RGBColor(0xF4, 0xC4, 0x30), False),
]


def shade_cell(cell, hex_color: str) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, *, size=11, bold=False, italic=False, color=INK, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def style_paragraph(p, *, space_before=0, space_after=8, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    return run


def heading(doc, text, level=1, center=False):
    p = doc.add_heading(text, level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(
            run,
            size=22 if level == 1 else 13 if level == 2 else 12,
            bold=True,
            color=INK,
        )
    style_paragraph(p, space_before=14 if level > 1 else 6, space_after=8)
    return p


def hrule(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), "D9D9D9")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=13, bold=True, color=TEAL)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run._r.append(instr)
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld2)


def apply_brand_chrome(section) -> None:
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = hp.add_run()
    run.add_picture(str(ASSETS / "header.png"), width=Inches(7.15))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)
    mp = footer.add_paragraph()
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mrun = mp.add_run()
    mrun.add_picture(str(ASSETS / "mascot.png"), width=Inches(0.34))


def body(doc, text, *, italic=False, color=INK, after=8, rule=False):
    p = doc.add_paragraph()
    style_paragraph(p, space_after=after)
    add_text(p, text, size=11, italic=italic, color=color)
    if rule:
        hrule(p)
    return p


def quote(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(p, space_before=12, space_after=6, line=1.3)
    add_text(p, f"“{text}”", size=13, italic=True)
    return p


def bullet(doc, text, level=0, bold=False):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    style_paragraph(p, space_before=1, space_after=2, line=1.12)
    add_text(p, text, size=11, bold=bold)
    return p


def role(doc, title, purpose, items):
    heading(doc, title, level=3)
    p = doc.add_paragraph()
    style_paragraph(p, space_after=4)
    add_text(p, "Core Purpose: ", size=11, bold=True)
    add_text(p, purpose, size=11)
    p = doc.add_paragraph()
    style_paragraph(p, space_before=4, space_after=4)
    add_text(p, "Responsibilities:", size=11, bold=True)
    for text, level in items:
        bullet(doc, text, level=level)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    for i in range(1, 4):
        h = doc.styles[f"Heading {i}"]
        h.font.name = FONT
        h.font.bold = True
        h.font.color.rgb = INK


def build() -> Path:
    import sys

    if str(ROOT) not in sys.path:
        sys.path.insert(0, str(ROOT))
    if not (ASSETS / "header.png").exists():
        from export_assets import export
        export()

    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(28)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(4)
    apply_brand_chrome(section)

    core = doc.core_properties
    core.title = "Unnatural Intelligence Club Charter"
    core.author = "Unnatural Intelligence"
    core.subject = (
        "Pursue untapped potential. Integrate AI into daily and professional "
        "workflows. Optimize tokens, keep the human."
    )

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, space_before=160, space_after=8)
    add_text(title, "Unnatural Intelligence", size=32, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(sub, space_after=10)
    add_text(sub, "Club Charter", size=22, bold=True)

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(date, space_after=14)
    add_text(date, "August 2026", size=13)

    tag = doc.add_paragraph()
    tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(tag, space_after=6)
    add_text(tag, "Untapped  ·  Integrated  ·  Human", size=11, italic=True, color=SLATE)

    doc.add_page_break()

    heading(doc, "Contents", center=True)
    toc = [
        ("Mission Statement", "3"),
        ("The Basis", "4"),
        ("Goals", "5"),
        ("Executive Roles & Responsibilities", "7"),
        ("Event Concepts", "13"),
        ("Style Guide", "16"),
    ]
    for name, num in toc:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(
            Mm(170), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
        )
        style_paragraph(p, space_after=6)
        add_text(p, name, size=12, bold=True)
        add_text(p, f"\t{num}", size=12, color=SLATE)

    doc.add_page_break()

    heading(doc, "Mission Statement", center=True)
    quote(
        doc,
        "Pursue the untapped potential of AI. "
        "Explore integrating AI into our daily and professional workflows. "
        "Optimize token usage while still retaining that touch of humanity.",
    )
    pillars = doc.add_paragraph()
    pillars.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(pillars, space_before=18, space_after=12)
    add_text(
        pillars,
        "Untapped potential  ·  Daily & professional workflows  ·  Lean tokens, human finish",
        size=10,
        color=SLATE,
    )

    doc.add_page_break()

    heading(doc, "The Basis", center=True)
    for title_text, para in [
        (
            "1. Pursue the untapped potential of AI",
            "Most students only meet the default chatbot. The interesting work is "
            "off to the side: multimodal shop-floor photos, syllabus-to-study-plan "
            "systems, custom assistants for one recurring task, uses nobody assigned "
            "in class. We go looking for those.",
        ),
        (
            "2. Integrate AI into daily and professional workflows",
            "A demo is not a habit. We treat AI as something you fold into Tuesday — "
            "class, practicum, job search, client email, shop estimates — then keep "
            "only the pieces that actually save time.",
        ),
        (
            "3. Optimize tokens, keep the human",
            "Wasteful prompts are sloppy thinking. We practice getting the same "
            "(or better) result with less context, then putting a human finish on "
            "whatever comes back so it still sounds like a person, not a template.",
        ),
        (
            "4. A polytechnic is the right place for this",
            "SAIT workflows are not generic. Construction, health, energy, business, "
            "media, and ICT each hide different unused potential. The club exists so "
            "those workflows can be compared in one room.",
        ),
        (
            "5. Built to exist as a Saitsa club",
            "Open membership, student-led, non-profit, three signing executives, "
            "registered through the Saitsa Clubs Portal. That structure is how we "
            "book rooms, apply for grants, and operate officially on campus.",
        ),
    ]:
        heading(doc, title_text, level=2)
        body(doc, para, rule=True, after=10)

    doc.add_page_break()

    heading(doc, "Goals", center=True)
    for title_text, para in [
        (
            "1. Pursue Untapped Potential",
            "Look past the default chatbot. Find uses, tools, and questions that "
            "students and workplaces have not worn out yet.",
        ),
        (
            "2. Integrate AI into Daily & Professional Workflows",
            "Move from one-off demos to habits: class, practicum, job search, "
            "and the work of each SAIT program.",
        ),
        (
            "3. Optimize Tokens, Keep the Human",
            "Spend fewer tokens for a better result, then edit until the work still "
            "sounds like a person — with judgment, taste, and a name on it.",
        ),
        (
            "4. Cross-Program Workflow Sharing",
            "A construction estimate, a lab write-up, and a client email are different "
            "jobs. Bring them into the same room so unused potential travels.",
        ),
        (
            "5. A Community of Practice, Not a Tool Fan Club",
            "Share the workflows people actually keep. Critique the ones that waste "
            "tokens or flatten a human voice.",
        ),
    ]:
        heading(doc, title_text, level=2)
        body(doc, para, rule=True, after=10)

    doc.add_page_break()

    for title_text, para in [
        (
            "6. Lower Barriers to Participation",
            "Make joining easy. You do not need a computer-science background, a paid "
            "model, or a finished prompt library to walk in.",
        ),
        (
            "7. Leadership Development",
            "Run the club the way we want people to run AI: lean, documented, and "
            "human. Meetings, budgets, events, and conflict — under Saitsa’s framework.",
        ),
        (
            "8. Experiment Without Waste",
            "Try the weird idea. Kill the prompt that burns tokens for nothing. "
            "Curiosity is encouraged; sloppy generation is not the culture.",
        ),
        (
            "9. Integrity, Voice, and Human Accountability",
            "Keep academic honesty and a recognizable human voice. If you could not "
            "stand behind the output without the model, it is not finished.",
        ),
        (
            "10. Strengthen Unnatural Intelligence’s Presence on Campus",
            "Be the SAIT club known for unused potential, real workflows, and work "
            "that still sounds like a person.",
        ),
    ]:
        heading(doc, title_text, level=2)
        body(doc, para, rule=True, after=10)

    doc.add_page_break()
    spacer = doc.add_paragraph()
    spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(spacer, space_before=200, space_after=4)
    add_text(spacer, "Executive Roles &", size=26, bold=True)
    spacer2 = doc.add_paragraph()
    spacer2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(spacer2, space_after=8)
    add_text(spacer2, "Responsibilities", size=26, bold=True)

    doc.add_page_break()

    heading(doc, "Directors")
    role(
        doc,
        "President — Strategic Lead & Club Coordinator",
        "Keep the club organized, aligned, and progressing forward.",
        [
            ("Hold veto power for decisions affecting club direction or budget.", 0),
            ("Set semester direction, themes, and goals around the three foundations.", 0),
            ("Arrange and lead executive meetings.", 0),
            ("Prepare agendas.", 1),
            ("Follow up with other execs on tasks.", 0),
            ("Track deadlines.", 1),
            ("Book rooms and manage logistics.", 0),
            ("Monitor club email, Teams, Rubric, and official communication channels.", 0),
            ("Navigate communication between execs during conflict.", 0),
            ("Represent Unnatural Intelligence externally, including with Saitsa.", 0),
            ("Ensure the three Saitsa signing authorities remain filled.", 0),
        ],
    )
    role(
        doc,
        "Vice-President Operations — Operations & Visionary",
        "Aid overall direction and maintain operations in the absence of the president.",
        [
            ("Serve as secondary decision-maker when the president is unavailable.", 0),
            ("Maintain direction of the club throughout the semester.", 0),
            ("Keep activity tied to the three foundations.", 1),
            ("Record meeting minutes.", 0),
            ("Assist the president with booking rooms and managing logistics.", 0),
            ("Maintain organization of documents.", 0),
            ("Event plans", 1),
            ("Budgets and grant paperwork", 1),
            ("Saitsa forms, Rubric submissions, and transition notes", 1),
            ("Help onboard new execs.", 0),
            ("Monitor club email, Teams, Rubric, and official communication channels.", 0),
            ("Navigate communication between execs during conflict.", 0),
            ("Represent Unnatural Intelligence externally.", 0),
        ],
    )
    role(
        doc,
        "Vice-President Finance — Budgeting & Funding",
        "Manage club finances responsibly and transparently.",
        [
            ("Prepare semester budget forecasts.", 0),
            ("Track overall expenses and maintain a ledger.", 0),
            ("Manage the Saitsa-issued club bank account.", 1),
            ("Keep receipts for the full club year for possible audit.", 1),
            ("Coordinate budgets with Events Coordinators.", 0),
            ("Take vision from an event and budget it out.", 1),
            ("Submit quotes and grant applications for events.", 1),
            ("Apply for new-club start-up reimbursement (up to $500) and term grants.", 0),
            ("Spend like we prompt: enough to do the work, nothing extra.", 0),
            ("Ensure purchases follow Saitsa and SAIT guidelines.", 0),
            ("Inform the SECC of any external monetary sponsorship.", 0),
            ("Reimburse members from the club account after approved spend.", 0),
        ],
    )
    body(
        doc,
        "Saitsa requires every club to keep three key executives filled: "
        "President, Vice-President Operations, and Vice-President Finance. "
        "These are the club’s signing authorities. Only current SAIT students "
        "who have paid their Saitsa fee may hold executive positions. "
        "Executives cannot be paid; the club operates on a non-profit basis.",
        italic=True,
        color=SLATE,
    )

    doc.add_page_break()
    heading(doc, "Marketing")
    role(
        doc,
        "Social Media Lead (1–3 positions) — Communication & Engagement",
        "Maintain active, engaging communication across platforms.",
        [
            ("Manage socials through regular, consistent posting and active replies.", 0),
            ("Instagram", 1),
            ("Discord", 1),
            ("Teams", 1),
            ("Rubric / Saitsa Clubs Portal", 1),
            ("LinkedIn", 1),
            ("Independently create captions, announcements, event recaps, stories, and engagement posts.", 0),
            ("In collaboration with Graphic Designers create general promotion.", 0),
            ("In collaboration with Graphic Designers and Event Coordination create event promotion.", 0),
            ("Track and optimize engagement metrics.", 0),
            ("Ensure consistent branding and a human club voice — no unedited AI captions.", 0),
            ("Publish lean-prompt tips, unused-use finds, and human-finish before/afters.", 0),
        ],
    )
    role(
        doc,
        "Graphic Designer (1–2 positions) — Visual Identity & Creative Production",
        "Shape the club’s visual identity and produce design assets.",
        [
            ("Design graphics for social media and any general or event-specific promotional material.", 0),
            ("Create templates for recurring content.", 0),
            ("Social media posts", 1),
            ("Posters (tabloid 11×17 for Saitsa boards, with the Saitsa Clubs logo)", 1),
            ("Presentations", 1),
            ("In collaboration with Social Media Leads create general promotion.", 0),
            ("In collaboration with Social Media Leads and Event Coordination create event promotion.", 0),
            ("Ensure consistent and distinct branding.", 0),
            ("Expand upon branding; test and experiment with designs.", 1),
            ("AI may sketch; a human always finishes what gets published.", 0),
            ("Do not use the official SAIT logo on club materials — student groups are not permitted to.", 0),
        ],
    )

    doc.add_page_break()
    heading(doc, "Events")
    role(
        doc,
        "Events Coordinator (2 positions) — Planning & Execution",
        "Plan and deliver high-quality events.",
        [
            ("Brainstorm and propose event ideas, including agenda, materials, and setup.", 0),
            ("Name which foundation the event serves (untapped, workflow, or human/token).", 0),
            ("Coordinate with Finance on event budgets.", 0),
            ("Coordinate with Marketing for promotional materials.", 0),
            ("Collaborate with Outreach to secure speakers.", 0),
            ("Submit event proposals through the Saitsa Clubs Portal before promoting.", 0),
            ("Follow up with Marketing on event tasks.", 0),
            ("Manage event logistics and oversee execution.", 0),
            ("Collect and compile feedback for future improvements.", 0),
        ],
    )
    role(
        doc,
        "Outreach & Networking — External Communication",
        "Connect the club with speakers, partners, and opportunities.",
        [
            ("Reach out to potential speakers, collaborators, or partnerships.", 0),
            ("Draft outreach messages", 1),
            ("Negotiate presentation details", 1),
            ("Maintain a contact list.", 0),
            ("Identify collaboration opportunities on campus and in Calgary industry.", 0),
            ("Find people whose real workflows still have unused AI potential.", 1),
            ("Coordinate with Events Coordinators to plan and arrange guests.", 0),
            ("In collaboration with Social Media Leads, manage LinkedIn connections.", 0),
            ("Represent Unnatural Intelligence externally.", 0),
            ("Help identify a faculty advisor (encouraged; required if registered as academic).", 0),
        ],
    )

    doc.add_page_break()
    spacer = doc.add_paragraph()
    spacer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(spacer, space_before=220)
    add_text(spacer, "Event Concepts", size=26, bold=True)

    doc.add_page_break()
    heading(doc, "Event Concepts")

    heading(doc, "Untapped Potential — find what default chat never gets asked.", level=2)
    for text, level in [
        ("“What’s this button for?” nights: one tool, the features nobody opens.", 0),
        ("Multimodal labs: a shop-floor photo, a schematic, a messy syllabus, a voice note.", 0),
        ("Build a tiny assistant for one recurring student task — then delete it if it is not used.", 0),
        ("Cross-program “I didn’t know it could do that” demos.", 0),
        ("Estimating, spec-checking, and photo-to-takeoff for construction and energy", 1),
        ("Charting, patient-education drafts, and protocol lookup for health", 1),
        ("Campaign research, spreadsheet cleanup, and client-ready copy for business", 1),
        ("Accessibility, translation, and study systems as unused campus uses.", 0),
        ("Failed-experiment night: the prompt that looked clever and did nothing.", 0),
    ]:
        bullet(doc, text, level=level)

    heading(doc, "Workflow Integration — from demo to a Tuesday you would actually keep.", level=2)
    for text, level in [
        ("Bring Your Tuesday: rebuild a real day (class, shift, job hunt) with AI in the loop.", 0),
        ("Syllabus to system: one course, one stack, one week of use.", 0),
        ("Professional pack: meeting notes, email, portfolios, interview drills.", 0),
        ("Two-week Daily Driver challenge: keep only the workflow you still open on day 14.", 0),
        ("Workflow swap: pair with a student from another program and steal one habit.", 0),
        ("Open studio (Prompt Playgrounds) for unfinished, in-progress work.", 0),
        ("Student-led: host a session on the workflow you actually use, not a tool tour.", 0),
        ("End of semester Daily Driver showcase — the habit that survived, not the flashiest demo.", 0),
    ]:
        bullet(doc, text, level=level)

    doc.add_page_break()

    heading(doc, "Token Craft & Human Touch — less spend, more of you in the result.", level=2)
    for text, level in [
        ("Token Diet: same outcome, fewer tokens. Publish the before/after prompt.", 0),
        ("Human Finish Studio: take a generic draft and edit until it sounds like you.", 0),
        ("“Don’t send that” clinic: catch the telltale AI voice before it hits an inbox.", 0),
        ("Context packing: what belongs in the prompt, and what you should leave out.", 0),
        ("When not to generate — academic integrity as keeping the human accountable.", 0),
        ("Blind taste tests: human, machine, and human-finished side by side.", 0),
        ("Social drops: five-minute lean-prompt tips, unused-use finds, human-finish of the week.", 0),
        ("Monthly challenge: one task, a token budget, a required human pass.", 0),
    ]:
        bullet(doc, text, level=level)

    heading(doc, "Mixer events", level=2)
    for text, level in [
        ("Game nights (phones down for a stretch — practice being the human in the room).", 0),
        ("Lean-token jams", 0),
        ("Unnatural-thons with a published brief and a token budget", 1),
        ("Short sprints that must ship a human-finished demo", 1),
        ("Workflow swap socials across programs", 0),
        ("Lunch & learns and industry nights with people who kept a workflow, not a slide deck", 0),
        ("Clubs Expo tables in fall and winter", 0),
    ]:
        bullet(doc, text, level=level)

    heading(doc, "How we register", level=2)
    for note_text in [
        "Any currently enrolled SAIT student can start or join. Before Stage 1 on the Saitsa Clubs Portal (Rubric), confirm three executives: President, Vice-President Operations, and Vice-President Finance.",
        "Register at portal.hellorubric.com, select Southern Alberta Institute of Technology, and submit for Saitsa approval. Contact saitsa.clubs@edu.sait.ca with questions.",
        "Membership stays open to all SAIT students. At least 75% of members must be current SAIT students. Community members may join at the executives’ discretion but cannot hold executive roles or vote.",
        "Seek a faculty advisor in the School for Advanced Digital Technology (or another school if a better fit). Advisors mentor; they do not run the club.",
    ]:
        bullet(doc, note_text)

    doc.add_page_break()
    heading(doc, "Unnatural Intelligence Style Guide", center=True)
    mark = doc.add_paragraph()
    mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(mark, space_before=8, space_after=12)
    mark.add_run().add_picture(str(ASSETS / "wordmark.png"), width=Inches(2.4))

    body(doc, "Google-friendly fonts: Inter  |  JetBrains Mono")
    body(
        doc,
        "Primary font: Inter (Inter Display for titles). Secondary font: JetBrains Mono for labels, codes, and UI moments.",
    )

    heading(doc, "Copy-paste colour codes", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for i, (hex_code, rgb, light_text) in enumerate(SWATCHES):
        cell = table.rows[i // 3].cells[i % 3]
        shade_cell(cell, hex_code)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, hex_code, size=11, bold=True, color=WHITE if light_text else INK, name="JetBrains Mono")
    # empty ninth cell
    empty = table.rows[2].cells[2]
    empty.text = ""

    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
