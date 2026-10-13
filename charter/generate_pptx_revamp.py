#!/usr/bin/env python3
"""A4 slide deck matching the 2026–27 branded charter."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

import revamp_copy as copy
from revamp_kit import ASSETS, export as export_kit

ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Unnatural_Intelligence_Club_Charter_2026-27.pptx"

INK = RGBColor(0x03, 0x02, 0x05)
WINE = RGBColor(0x59, 0x13, 0x37)
SLATE = RGBColor(0x5A, 0x5A, 0x5A)
PAPER = RGBColor(0xF9, 0xF9, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PINK = RGBColor(0xFA, 0x4A, 0x9E)
OLIVE = RGBColor(0x95, 0xB7, 0x47)
LIME = RGBColor(0xBA, 0xF8, 0x37)
CYAN = RGBColor(0x7E, 0xC8, 0xFF)
RULE = RGBColor(0xD7, 0xD7, 0xDA)

SWATCHES = [
    ("#F9F9FA", PAPER, False),
    ("#95B747", OLIVE, False),
    ("#030205", INK, True),
    ("#FA4A9E", PINK, True),
    ("#591337", WINE, True),
    ("#BAF837", LIME, False),
    ("#7EC8FF", CYAN, False),
]


def _prs() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(8.27)
    prs.slide_height = Inches(11.69)
    return prs


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def fill_slide(slide, color=PAPER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(8.27), Inches(11.69)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def chrome(slide, prs, page: int) -> None:
    fill_slide(slide)
    slide.shapes.add_picture(
        str(ASSETS / "header.png"),
        Inches(0.16),
        Inches(0.08),
        width=prs.slide_width - Inches(0.32),
    )
    box = slide.shapes.add_textbox(Inches(0), Inches(11.18), prs.slide_width, Inches(0.22))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"{page:02d}"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = WINE
    run.font.name = "JetBrains Mono"
    slide.shapes.add_picture(
        str(ASSETS / "jellyfish-mini.png"),
        prs.slide_width / 2 - Inches(0.18),
        Inches(11.28),
        width=Inches(0.36),
    )


def tb(slide, l, t, w, h, text, *, size=12, bold=False, italic=False, color=INK, align=PP_ALIGN.LEFT, font="Inter"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    parts = str(text).split("\n")
    for i, part in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = part
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    return tf


def bullets(slide, l, t, w, h, items, size=13):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, level) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = min(level, 2)
        p.font.size = Pt(size if level == 0 else size - 1)
        p.font.color.rgb = INK
        p.font.name = "Inter"
        p.space_after = Pt(4)
    return tf


def section_list(slide, prs, title, blocks, page):
    chrome(slide, prs, page)
    tb(slide, Inches(0.55), Inches(1.15), Inches(7.2), Inches(0.45), title, size=24, bold=True, align=PP_ALIGN.CENTER)
    y = 1.7
    for heading, body in blocks:
        tb(slide, Inches(0.7), Inches(y), Inches(6.9), Inches(0.32), heading, size=14, bold=True)
        tb(slide, Inches(0.7), Inches(y + 0.30), Inches(6.9), Inches(0.85), body, size=12)
        y += 1.22


def build() -> Path:
    if not (ASSETS / "header.png").exists():
        export_kit()

    prs = _prs()
    M = Inches(0.7)
    W = Inches(6.9)
    page = 1

    s = blank(prs)
    chrome(s, prs, page)
    s.shapes.add_picture(str(ASSETS / "jellyfish.png"), Inches(3.05), Inches(1.35), height=Inches(2.4))
    s.shapes.add_picture(str(ASSETS / "wordmark.png"), Inches(1.9), Inches(3.85), width=Inches(4.5))
    tb(s, Inches(0.4), Inches(5.55), Inches(7.5), Inches(0.5), "Club Charter", size=28, bold=True, align=PP_ALIGN.CENTER)
    tb(s, Inches(0.4), Inches(6.1), Inches(7.5), Inches(0.35), copy.YEAR, size=14, bold=True, color=WINE, align=PP_ALIGN.CENTER, font="JetBrains Mono")
    tb(s, Inches(0.4), Inches(6.5), Inches(7.5), Inches(0.35), copy.PILLARS, size=12, italic=True, color=SLATE, align=PP_ALIGN.CENTER)
    s.shapes.add_picture(str(ASSETS / "tagline.png"), Inches(2.4), Inches(7.05), width=Inches(3.5))
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    tb(s, M, Inches(1.2), W, Inches(0.4), "Contents", size=22, bold=True, align=PP_ALIGN.CENTER)
    bullets(s, M, Inches(1.9), W, Inches(7), [(t, 0) for t, _n, bold, _l in copy.TOC if bold], size=16)
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    tb(s, M, Inches(1.3), W, Inches(0.45), "Mission Statement", size=24, bold=True, align=PP_ALIGN.CENTER)
    tb(s, Inches(0.75), Inches(1.85), Inches(6.8), Inches(1.15), copy.MISSION_LEAD, size=13, align=PP_ALIGN.CENTER)
    tb(s, Inches(0.75), Inches(3.05), Inches(6.8), Inches(1.25), f"“{copy.MISSION_QUOTE}”", size=13, italic=True, color=WINE, align=PP_ALIGN.CENTER)
    tb(s, Inches(0.75), Inches(4.35), Inches(6.8), Inches(0.95), copy.MISSION_FEELINGS, size=13, align=PP_ALIGN.CENTER)
    card_w = Inches(2.12)
    gap = Inches(0.18)
    total = card_w * 3 + gap * 2
    start = (prs.slide_width - total) / 2
    card_top = Inches(5.45)
    card_h = Inches(2.35)
    icon_h = Inches(0.40)
    for i, (icon, title, body) in enumerate(copy.FOUNDATIONS):
        x = start + i * (card_w + gap)
        shape = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x, card_top, card_w, card_h
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = RULE
        pad = Inches(0.18)
        s.shapes.add_picture(
            str(ASSETS / f"icon-{icon}.png"),
            x + (card_w - icon_h) / 2,
            card_top + pad,
            height=icon_h,
        )
        tb(
            s,
            x + pad,
            card_top + pad + icon_h + Inches(0.18),
            card_w - pad * 2,
            Inches(0.50),
            title,
            size=11,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        tb(
            s,
            x + pad,
            card_top + pad + icon_h + Inches(0.76),
            card_w - pad * 2,
            Inches(1.10),
            body,
            size=10,
            color=SLATE,
            align=PP_ALIGN.CENTER,
        )
    s.shapes.add_picture(str(ASSETS / "quote-useful.png"), Inches(2.55), Inches(8.05), width=Inches(3.15))
    page += 1

    s = blank(prs)
    section_list(s, prs, "The Basis", copy.BASIS, page)
    page += 1

    s = blank(prs)
    section_list(s, prs, "Goals", copy.GOALS_A, page)
    page += 1
    s = blank(prs)
    section_list(s, prs, "Goals", copy.GOALS_B, page)
    page += 1

    s = blank(prs)
    section_list(s, prs, "How We Operate", copy.OPERATE, page)
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    s.shapes.add_picture(str(ASSETS / "jellyfish.png"), Inches(3.15), Inches(3.2), height=Inches(2.2))
    tb(s, M, Inches(5.6), W, Inches(0.5), "Executive Roles &", size=28, bold=True, align=PP_ALIGN.CENTER)
    tb(s, M, Inches(6.15), W, Inches(0.5), "Responsibilities", size=28, bold=True, align=PP_ALIGN.CENTER)
    page += 1

    def role_slide(group, role_tuple, pg):
        sl = blank(prs)
        chrome(sl, prs, pg)
        title, purpose, items = role_tuple
        tb(sl, M, Inches(1.15), W, Inches(0.35), group, size=20, bold=True)
        tb(sl, M, Inches(1.55), W, Inches(0.45), title, size=14, bold=True)
        tb(sl, M, Inches(2.05), W, Inches(0.4), f"Core Purpose: {purpose}", size=12)
        tb(sl, M, Inches(2.45), W, Inches(0.3), "Responsibilities:", size=12, bold=True)
        bullets(sl, M, Inches(2.8), W, Inches(7.6), items, size=13)
        return sl

    role_slide("Directors", copy.PRESIDENT, page)
    page += 1
    role_slide("Directors", copy.VP_OPS, page)
    page += 1
    role_slide("Directors", copy.VP_FINANCE, page)
    page += 1
    role_slide("Marketing", copy.SOCIAL, page)
    page += 1
    role_slide("Marketing", copy.DESIGNER, page)
    page += 1
    role_slide("Events", copy.EVENTS_COORD, page)
    page += 1
    role_slide("Events", copy.OUTREACH, page)
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    tb(s, M, Inches(5.2), W, Inches(0.5), "Event Concepts", size=28, bold=True, align=PP_ALIGN.CENTER)
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    tb(s, M, Inches(1.2), W, Inches(0.5), "Untapped Potential", size=16, bold=True)
    bullets(s, M, Inches(1.7), W, Inches(3.6), [(t, max(lv - 1, 0)) for t, lv in copy.UNTAPPED_ITEMS], size=13)
    tb(s, M, Inches(5.6), W, Inches(0.5), "Workflow Integration", size=16, bold=True)
    bullets(s, M, Inches(6.15), W, Inches(4.2), [(t, max(lv - 1, 0)) for t, lv in copy.WORKFLOW_ITEMS], size=13)
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    tb(s, M, Inches(1.2), W, Inches(0.5), "Token Craft & Human Touch", size=16, bold=True)
    bullets(s, M, Inches(1.7), W, Inches(4.4), [(t, max(lv - 1, 0)) for t, lv in copy.TOKEN_ITEMS], size=13)
    tb(s, M, Inches(6.4), W, Inches(0.4), "Mixer events", size=16, bold=True)
    bullets(s, M, Inches(6.85), W, Inches(3.6), [(t, max(lv - 1, 0)) for t, lv in copy.MIXER_ITEMS], size=13)
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    tb(s, M, Inches(1.2), W, Inches(0.4), "This Year on Campus", size=22, bold=True, align=PP_ALIGN.CENTER)
    bullets(
        s,
        M,
        Inches(1.8),
        W,
        Inches(5.5),
        [(f"{date} — {event}", 0) for date, event in copy.CALENDAR],
        size=14,
    )
    page += 1

    s = blank(prs)
    chrome(s, prs, page)
    tb(s, M, Inches(1.15), W, Inches(0.4), "Unnatural Intelligence Style Guide", size=20, bold=True, align=PP_ALIGN.CENTER)
    s.shapes.add_picture(str(ASSETS / "jellyfish.png"), Inches(3.15), Inches(1.7), height=Inches(1.8))
    top = Inches(3.8)
    box_w, box_h = Inches(1.55), Inches(0.72)
    gap = Inches(0.18)
    n = len(SWATCHES)
    for i, (hex_code, rgb, light) in enumerate(SWATCHES):
        col, row = i % 4, i // 4
        items_in_row = min(4, n - row * 4)
        row_width = items_in_row * box_w + (items_in_row - 1) * gap
        row_left = (prs.slide_width - row_width) / 2
        shape = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            row_left + col * (box_w + gap),
            top + row * (box_h + Inches(0.28)),
            box_w,
            box_h,
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb
        shape.line.fill.background()
        tf = shape.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = hex_code
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = WHITE if light else INK
        run.font.name = "JetBrains Mono"

    prs.save(str(OUTPUT))
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
